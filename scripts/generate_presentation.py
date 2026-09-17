"""
Generates the EXACTLY 10-slide PowerPoint presentation for Assignment 3:
AirResolve - Customer-Facing Resolution Agent (Airline Disruption).
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY = RGBColor(11, 25, 44)       # #0B192C
    C_BLUE = RGBColor(30, 62, 98)       # #1E3E62
    C_ACCENT = RGBColor(0, 141, 218)    # #008DDA
    C_BG = RGBColor(248, 250, 252)      # #F8FAFC
    C_WHITE = RGBColor(255, 255, 255)
    C_DARK = RGBColor(15, 23, 42)       # #0F172A
    C_MUTED = RGBColor(100, 116, 139)   # #64748B
    C_CARD_BG = RGBColor(255, 255, 255)
    C_BORDER = RGBColor(226, 232, 240)
    C_GREEN = RGBColor(5, 150, 105)
    C_RED = RGBColor(220, 38, 38)
    C_AMBER = RGBColor(217, 119, 6)

    def add_header(slide, title_text, category_text="AIRRESOLVE • DISRUPTION RESOLUTION AGENT"):
        # Header banner shape
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.1))
        shape.fill.solid()
        shape.fill.fore_color.rgb = C_NAVY
        shape.line.color.rgb = C_NAVY

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.4)
        tf.margin_top = Inches(0.18)

        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = C_ACCENT

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

    def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.2)
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    # Background rectangle
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY
    bg1.line.fill.background()

    # Title card inside
    card1 = add_card(s1, 1.2, 1.2, 10.933, 5.1, bg_color=C_BLUE, border_color=C_ACCENT)
    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.8)
    tf1.margin_top = Inches(0.8)

    p = tf1.paragraphs[0]
    p.text = "AIRRESOLVE: AIRLINE DISRUPTION RESOLUTION AGENT"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p = tf1.add_paragraph()
    p.text = "Individual Assignment 3: Customer-Facing Resolution Agent"
    p.font.size = Pt(18)
    p.font.color.rgb = C_ACCENT
    p.space_before = Pt(14)

    p = tf1.add_paragraph()
    p.text = "A policy-grounded conversational agent with deterministic business rules, auditability, and guardrails."
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(203, 213, 225)
    p.space_before = Pt(10)

    p = tf1.add_paragraph()
    p.text = "Candidate / Student Name: [PASTE YOUR NAME HERE]\nEvaluated Against: Assignment 3 Ground Truth Data Pack (Date: 23 September 2026)"
    p.font.size = Pt(12.5)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_before = Pt(35)

    # =========================================================================
    # SLIDE 2: Problem Statement
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Airline Disruption Challenge & Automated Resolution", "Phase 1 • Problem Context")

    # Card Left: Problem
    c2_left = add_card(s2, 0.8, 1.9, 5.7, 4.9)
    tf2_l = c2_left.text_frame
    tf2_l.word_wrap = True
    tf2_l.margin_left = Inches(0.4)
    tf2_l.margin_top = Inches(0.4)

    p = tf2_l.paragraphs[0]
    p.text = "Current Disruption Pain Points"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_RED

    points_l = [
        "High Friction & Long Wait Times: Flight cancellations and multi-hour delays cause extreme passenger surge at airport desks and contact centers.",
        "Unpredictable Agent Hallucinations: Standard generative LLMs hallucinate policies, promising unwarranted upgrades or illegal refunds.",
        "Compliance & Financial Leakage: Frontline errors in waiving fare differences or approving unauthorized hotel stays lead to major revenue loss.",
        "Lack of Auditability: Traditional chatbots fail to produce clear compliance traces explaining why a compensation benefit was granted or denied."
    ]
    for pt in points_l:
        p = tf2_l.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(10)

    # Card Right: Solution
    c2_right = add_card(s2, 6.833, 1.9, 5.7, 4.9)
    tf2_r = c2_right.text_frame
    tf2_r.word_wrap = True
    tf2_r.margin_left = Inches(0.4)
    tf2_r.margin_top = Inches(0.4)

    p = tf2_r.paragraphs[0]
    p.text = "AirResolve Solution Architecture"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    points_r = [
        "Deterministic Policy Engine: Business rules are codified in pure Python logic as the single, uncompromised source of truth.",
        "Empathetic Yet Gated AI: The conversational layer understands natural language and user frustration, but cannot override policies.",
        "Transparent Audit Trail: Every response includes a complete decision trace showing customer, booking status, and policy evaluation steps.",
        "Proactive Guardrails & Escalation: Automatically escalates legal threats, complaints, and waivers exceeding the ₹1,500 agent threshold."
    ]
    for pt in points_r:
        p = tf2_r.add_paragraph()
        p.text = f"✔ {pt}"
        p.font.size = Pt(12)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 3: Assignment Data & Scenarios
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Assignment 3 Data Pack Ground Truth & Profiles", "Phase 2 • Operational Context (23 Sep 2026)")

    # 3 Scenario Cards
    scenarios_data = [
        ("Priya Nair", "Gold", "SK4821X", "Delhi → Goa (SK-204)", "Cancelled", C_RED,
         "Operational reasons. Return flight (25 Sep) unaffected.",
         "Requests full cash refund + complimentary Business Class upgrade on return flight."),
        ("Arvind Kulkarni", "Silver", "TR1190B", "Mumbai → Bengaluru (SK-118)", "Delayed 4h", C_AMBER,
         "New departure: 11:10 (Original: 07:10). Missed connecting meeting.",
         "Requests hotel accommodation for the 4-hour delay."),
        ("Meher Kaur", "Platinum", "WL7742", "Delhi → Hyderabad (SK-305)", "Delayed 6h", C_BLUE,
         "New departure: 20:00 (Original: 14:00). High-tier flyer (10 flights).",
         "Requests full night's hotel stay + voluntary rebooking on higher-fare flight (₹2,000 diff)."),
    ]

    left_pos = 0.8
    for name, tier, pnr, route, status, color, context, req in scenarios_data:
        c = add_card(s3, left_pos, 1.9, 3.65, 4.9)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = C_NAVY

        p = tf.add_paragraph()
        p.text = f"Tier: {tier}  |  PNR: {pnr}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_ACCENT

        p = tf.add_paragraph()
        p.text = f"Route: {route}"
        p.font.size = Pt(11)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(6)

        p = tf.add_paragraph()
        p.text = f"Disruption: {status}"
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = f"Context: {context}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(6)

        p = tf.add_paragraph()
        p.text = "Passenger Request:"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_DARK
        p.space_before = Pt(12)

        p = tf.add_paragraph()
        p.text = f"\"{req}\""
        p.font.size = Pt(10.5)
        p.font.italic = True
        p.font.color.rgb = C_BLUE

        left_pos += 4.04

    # =========================================================================
    # SLIDE 4: Business / Service Rules
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Deterministic Business & Service Rules", "Phase 3 • Policy Framework")

    # Table or 4 Rule blocks
    rules_data = [
        ("Cancellation Rebooking Rule", "If flight is cancelled by airline, customer entitled to free rebooking within 24h OR full refund (customer's choice).", C_NAVY),
        ("Delay Compensation Rule", "• < 3h: ₹500 meal voucher\n• 3h-5h: ₹500 meal voucher + lounge access\n• > 5h: Meal voucher + lounge + hotel covering ONLY delayed hours.", C_BLUE),
        ("Refund Processing Rule", "Refunds for airline-caused disruptions are processed in full within 7 business days to original payment method only.", C_NAVY),
        ("Fare Difference & Loyalty Rules", "• Voluntary rebooking: Fare diff > ₹1,500 requires supervisor approval (strictly prohibited for agent).\n• Gold/Platinum: Priority rebooking only; NO additional compensation.", C_BLUE),
    ]

    positions = [(0.8, 1.9), (6.833, 1.9), (0.8, 4.45), (6.833, 4.45)]
    for i, (title, content, color) in enumerate(rules_data):
        l, t = positions[i]
        c = add_card(s4, l, t, 5.7, 2.35)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.text = content
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 5: System Architecture
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Complete AirResolve End-to-End Architecture", "Phase 4 • System Architecture")

    # Embed architecture diagram if it exists
    arch_img_path = os.path.join(os.path.dirname(__file__), "..", "assets", "architecture.png")
    if os.path.exists(arch_img_path):
        s5.shapes.add_picture(arch_img_path, Inches(0.8), Inches(1.8), width=Inches(8.5))

    # Right side architectural highlights card
    c5_r = add_card(s5, 9.5, 1.8, 3.033, 5.0)
    tf5_r = c5_r.text_frame
    tf5_r.word_wrap = True
    tf5_r.margin_left = Inches(0.25)
    tf5_r.margin_top = Inches(0.3)

    p = tf5_r.paragraphs[0]
    p.text = "Key Principles"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    principles = [
        ("Deterministic Gate:", "Policy engine is pure Python logic; cannot be bypassed."),
        ("No Policy Override:", "LLM extracts intent but cannot relax or invent rules."),
        ("Strict Thresholds:", "Hotel strictly > 5h; waiver ceiling strictly ₹1,500."),
        ("Structured Decisions:", "Every output yields a structured audit record with rule ID."),
    ]
    for title, desc in principles:
        p = tf5_r.add_paragraph()
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_ACCENT
        p.space_before = Pt(8)

        p = tf5_r.add_paragraph()
        p.text = desc
        p.font.size = Pt(10)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 6: Agent Decision Flow
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Decision Flow & Policy Evaluation Pipeline", "Phase 5 • Decision Engineering")

    flow_steps = [
        ("1. Passenger Identification", "Matches PNR or Name to customer profile. Retrieves loyalty tier and past travel/complaint history.", C_NAVY),
        ("2. Booking Verification", "Checks flight segment status (Cancelled, Delayed 4h, Delayed 6h). Validates airline-caused disruption.", C_BLUE),
        ("3. Intent & Entity Extraction", "Identifies primary intent (Refund, Hotel, Upgrade, Rebook, Fare diff) and amounts (e.g. ₹2,000, full-night).", C_NAVY),
        ("4. Policy Engine Evaluation", "Executes deterministic Python rule functions against parameters. Identifies allowed vs prohibited actions.", C_BLUE),
        ("5. Guardrail & Decision Gate", "RESOLVE: Executes refund/vouchers.\nCLARIFY: Prompts customer choice.\nESCALATE: Routes to supervisor.", C_ACCENT),
    ]

    for idx, (title, desc, color) in enumerate(flow_steps):
        top_pos = 1.9 + idx * 0.98
        c = add_card(s6, 0.8, top_pos, 11.733, 0.85)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_DARK

    # =========================================================================
    # SLIDE 7: Scenario 1 — Priya Nair
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Scenario 1: Priya Nair (Gold, SK4821X)", "Phase 6 • Scenario Verification")

    c7_l = add_card(s7, 0.8, 1.9, 5.7, 4.9)
    tf7_l = c7_l.text_frame
    tf7_l.word_wrap = True
    tf7_l.margin_left = Inches(0.4)
    tf7_l.margin_top = Inches(0.4)

    p = tf7_l.paragraphs[0]
    p.text = "Disruption Context & Customer Request"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf7_l.add_paragraph()
    p.text = "• Flight: SK-204 (Delhi → Goa) Cancelled for operational reasons.\n• Return: SK-205 (Goa → Delhi, 25 Sep) is Unaffected.\n• Customer State: Furious.\n• Requests: Full cash refund + complimentary upgrade to Business Class on return flight 'for the trouble'."
    p.font.size = Pt(12)
    p.font.color.rgb = C_DARK
    p.space_before = Pt(10)

    p = tf7_l.add_paragraph()
    p.text = "Immediate Escalation Trigger Test:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_RED
    p.space_before = Pt(20)

    p = tf7_l.add_paragraph()
    p.text = "If Priya threatens legal action or files a formal complaint, the agent immediately halts automated processing and triggers specialist escalation."
    p.font.size = Pt(11)
    p.font.color.rgb = C_DARK

    c7_r = add_card(s7, 6.833, 1.9, 5.7, 4.9)
    tf7_r = c7_r.text_frame
    tf7_r.word_wrap = True
    tf7_r.margin_left = Inches(0.4)
    tf7_r.margin_top = Inches(0.4)

    p = tf7_r.paragraphs[0]
    p.text = "Agent Behavior & Policy Rationale"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    p = tf7_r.add_paragraph()
    p.text = "1. Full Refund Approved (RESOLVED):\n   Processed in full within 7 business days to original payment method per Refund Processing Rule.\n\n2. Free Business Upgrade REFUSED (INELIGIBLE):\n   Policy strictly does NOT provide complimentary cabin upgrades for flight disruptions or loyalty tiers. Gold tier grants priority rebooking, not free upgrades.\n\n3. Empathetic Tone:\n   Validates frustration while upholding absolute compliance with policy boundaries."
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_DARK
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 8: Scenario 2 — Arvind Kulkarni
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Scenario 2: Arvind Kulkarni (Silver, TR1190B)", "Phase 6 • Scenario Verification")

    c8_l = add_card(s8, 0.8, 1.9, 5.7, 4.9)
    tf8_l = c8_l.text_frame
    tf8_l.word_wrap = True
    tf8_l.margin_left = Inches(0.4)
    tf8_l.margin_top = Inches(0.4)

    p = tf8_l.paragraphs[0]
    p.text = "Disruption Context & Customer Request"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf8_l.add_paragraph()
    p.text = "• Flight: SK-118 (Mumbai → Bengaluru).\n• Status: Delayed 4 hours (New departure: 11:10, Original: 07:10).\n• Customer State: Frustrated about missing a connecting meeting.\n• Request: Demands hotel accommodation 'since it's been such a long delay'."
    p.font.size = Pt(12)
    p.font.color.rgb = C_DARK
    p.space_before = Pt(10)

    c8_r = add_card(s8, 6.833, 1.9, 5.7, 4.9)
    tf8_r = c8_r.text_frame
    tf8_r.word_wrap = True
    tf8_r.margin_left = Inches(0.4)
    tf8_r.margin_top = Inches(0.4)

    p = tf8_r.paragraphs[0]
    p.text = "Agent Behavior & Policy Rationale"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    p = tf8_r.add_paragraph()
    p.text = "1. Delay Benefits Issued (RESOLVED):\n   4-hour delay (> 3 hours, <= 5 hours) entitles customer to ₹500 meal voucher and lounge access. Both are applied automatically.\n\n2. Hotel Accommodation DENIED (INELIGIBLE):\n   Delay Compensation Rule strictly states hotel accommodation applies ONLY to delays exceeding 5 hours. At 4 hours, Arvind is not eligible.\n\n3. Zero False Promises:\n   Agent explains the policy boundary transparently without capitulating to pressure."
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_DARK
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 9: Scenario 3 — Meher Kaur
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Scenario 3: Meher Kaur (Platinum, WL7742)", "Phase 6 • Scenario Verification")

    c9_l = add_card(s9, 0.8, 1.9, 5.7, 4.9)
    tf9_l = c9_l.text_frame
    tf9_l.word_wrap = True
    tf9_l.margin_left = Inches(0.4)
    tf9_l.margin_top = Inches(0.4)

    p = tf9_l.paragraphs[0]
    p.text = "Disruption Context & Customer Requests"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    p = tf9_l.add_paragraph()
    p.text = "• Flight: SK-305 (Delhi → Hyderabad).\n• Status: Delayed 6 hours (New departure: 20:00, Original: 14:00).\n• Customer: Platinum Tier frequent flyer (10 flights past year).\n• Request 1: Full night's hotel stay (rather than coverage for just delayed hours).\n• Request 2: Voluntary rebooking to different, higher-fare flight with ₹2,000 fare difference."
    p.font.size = Pt(12)
    p.font.color.rgb = C_DARK
    p.space_before = Pt(10)

    c9_r = add_card(s9, 6.833, 1.9, 5.7, 4.9)
    tf9_r = c9_r.text_frame
    tf9_r.word_wrap = True
    tf9_r.margin_left = Inches(0.4)
    tf9_r.margin_top = Inches(0.4)

    p = tf9_r.paragraphs[0]
    p.text = "Agent Behavior & Escalation Rationale"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_RED

    p = tf9_r.add_paragraph()
    p.text = "1. Delay Benefits & Partial Hotel (RESOLVED):\n   6-hour delay (> 5h) qualifies for ₹500 meal voucher, lounge access, and hotel accommodation covering ONLY delayed hours (until 20:00 departure). Full night stay is refused as it exceeds policy scope.\n\n2. Mandatory Escalation on ₹2,000 Fare Diff (ESCALATE):\n   Fare Difference Rule: Agents CANNOT waive fare differences above ₹1,500 without supervisor approval.\n   Because ₹2,000 > ₹1,500, the agent creates an Escalation Ticket (ESC-WL7742-XXXX) for supervisor review."
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_DARK
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 10: Implementation + Testing + Demo
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Implementation, Testing, Demo & GitHub Readiness", "Phase 7 • Deliverables & Verification")

    deliverables = [
        ("Tech Stack", "• Python 3.11\n• Streamlit (Modern Custom UI)\n• Deterministic Policy Engine\n• Pytest (100% Passing)\n• python-pptx / Matplotlib"),
        ("Test Suite", "• 28 Automated Tests\n• 100% scenario coverage\n• Cancellation, delay brackets\n• ₹1,500 waiver boundary\n• Legal threat escalation"),
        ("Deliverables", "• Streamlit App (app.py)\n• Exact 10-Slide PPT\n• Architecture Diagram\n• Demo Script (3-5 min)\n• Technical Defense Guide"),
        ("Demo & GitHub", "• GitHub: github.com/USER/REPO\n• Public Demo Video: Google Drive link placeholder\n• One-Command Local Run:\n  streamlit run app.py"),
    ]

    positions_10 = [(0.8, 1.9), (3.84, 1.9), (6.88, 1.9), (9.92, 1.9)]
    for i, (col_title, col_content) in enumerate(deliverables):
        l, t = positions_10[i]
        c = add_card(s10, l, t, 2.65, 4.9)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = col_title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = C_NAVY

        p = tf.add_paragraph()
        p.text = col_content
        p.font.size = Pt(11)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(10)

    out_deck_path = os.path.join(os.path.dirname(__file__), "..", "Assignment_3_10_Slide_PPT.pptx")
    prs.save(out_deck_path)
    print(f"PowerPoint successfully created with {len(prs.slides)} slides at: {out_deck_path}")

if __name__ == "__main__":
    create_deck()
